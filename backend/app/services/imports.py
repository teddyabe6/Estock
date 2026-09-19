"""Spreadsheet product import: upload → map → preview → confirm (PRD 8.3).

Nothing is written until the whole file has been validated, so an import is
never partial and unexplained.  Opening stock creates real stock movements.
"""

from __future__ import annotations

import csv
import io
import json
import uuid
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.audit import AuditAction, record_audit
from app.core.db import utcnow
from app.core.errors import ConflictError, ValidationError
from app.core.permissions import Permission
from app.core.tenancy import AuthContext, get_tenant_object, tenant_query
from app.models.catalogue import Product, ProductVariant
from app.models.system import FileAsset, ImportJob, ImportJobStatus, ImportRowError
from app.services.catalogue import ProductInput, create_product
from app.services.storage import (
    ALLOWED_IMPORT_TYPES,
    get_storage,
    validate_upload,
)

#: Fields an uploaded column can be mapped to, and whether they are required.
IMPORTABLE_FIELDS: dict[str, dict] = {
    "name": {"label": "Product name", "required": True},
    "sku": {"label": "SKU / code", "required": False},
    "barcode": {"label": "Barcode", "required": False},
    "category_name": {"label": "Category", "required": False},
    "brand": {"label": "Brand", "required": False},
    "unit_of_measure": {"label": "Unit", "required": False},
    "description": {"label": "Description", "required": False},
    "purchase_price": {"label": "Purchase price", "required": False, "numeric": True},
    "transport_cost": {"label": "Transport cost", "required": False, "numeric": True},
    "other_costs": {"label": "Other costs", "required": False, "numeric": True},
    "selling_price": {"label": "Selling price", "required": False, "numeric": True},
    "opening_stock": {"label": "Quantity in stock", "required": False, "numeric": True},
    "min_stock": {"label": "Minimum stock", "required": False, "numeric": True},
    "reorder_level": {"label": "Reorder level", "required": False, "numeric": True},
}

#: Header text the mapper recognises without the user choosing (PRD 8.3).
HEADER_SYNONYMS: dict[str, str] = {
    "product": "name",
    "product name": "name",
    "item": "name",
    "item name": "name",
    "name": "name",
    "sku": "sku",
    "code": "sku",
    "product code": "sku",
    "barcode": "barcode",
    "bar code": "barcode",
    "category": "category_name",
    "brand": "brand",
    "unit": "unit_of_measure",
    "uom": "unit_of_measure",
    "description": "description",
    "cost": "purchase_price",
    "purchase price": "purchase_price",
    "buying price": "purchase_price",
    "unit cost": "purchase_price",
    "transport": "transport_cost",
    "transport cost": "transport_cost",
    "other cost": "other_costs",
    "other costs": "other_costs",
    "price": "selling_price",
    "selling price": "selling_price",
    "sale price": "selling_price",
    "quantity": "opening_stock",
    "qty": "opening_stock",
    "stock": "opening_stock",
    "opening stock": "opening_stock",
    "minimum stock": "min_stock",
    "min stock": "min_stock",
    "reorder level": "reorder_level",
}

MAX_IMPORT_ROWS = 5000


@dataclass(slots=True)
class ParsedRow:
    row_number: int
    values: dict[str, str]


@dataclass(slots=True)
class RowIssue:
    row_number: int
    column: str | None
    code: str
    message: str
    raw_value: str | None = None


@dataclass(slots=True)
class ImportPreview:
    job_id: uuid.UUID
    total_rows: int
    valid_rows: int
    error_rows: int
    duplicate_rows: int
    sample: list[dict] = field(default_factory=list)
    issues: list[RowIssue] = field(default_factory=list)


def parse_spreadsheet(content: bytes, filename: str) -> tuple[list[str], list[ParsedRow]]:
    """Read an .xlsx or .csv file into a header row and data rows."""
    if filename.lower().endswith((".csv", ".txt")):
        return _parse_csv(content)
    return _parse_xlsx(content)


def _parse_csv(content: bytes) -> tuple[list[str], list[ParsedRow]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise ValidationError("The file has no rows")
    headers = [str(cell).strip() for cell in rows[0]]
    parsed = [
        ParsedRow(
            row_number=index + 2,
            values={
                headers[i]: str(cell).strip()
                for i, cell in enumerate(row)
                if i < len(headers) and str(cell).strip() != ""
            },
        )
        for index, row in enumerate(rows[1:])
        if any(str(cell).strip() for cell in row)
    ]
    return headers, parsed


def _parse_xlsx(content: bytes) -> tuple[list[str], list[ParsedRow]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surfaced to the user as a validation error
        raise ValidationError(
            "This file could not be read as a spreadsheet. Save it as .xlsx or .csv "
            "and try again."
        ) from exc
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise ValidationError("The file has no rows") from exc

    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
    parsed: list[ParsedRow] = []
    for index, row in enumerate(rows):
        if row is None or not any(cell is not None and str(cell).strip() for cell in row):
            continue
        values = {}
        for i, cell in enumerate(row):
            if i >= len(headers) or not headers[i] or cell is None:
                continue
            text = str(cell).strip()
            if text:
                values[headers[i]] = text
        parsed.append(ParsedRow(row_number=index + 2, values=values))
    workbook.close()
    return headers, parsed


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    """Best-guess column mapping the user can correct (PRD 8.3)."""
    mapping: dict[str, str] = {}
    for header in headers:
        key = header.strip().lower()
        field_name = HEADER_SYNONYMS.get(key)
        if field_name and field_name not in mapping.values():
            mapping[header] = field_name
    return mapping


def start_import(
    db: Session,
    ctx: AuthContext,
    *,
    filename: str,
    content_type: str,
    content: bytes,
    branch_id: uuid.UUID | None = None,
) -> tuple[ImportJob, list[str], dict[str, str]]:
    """Store the upload, detect headers and suggest a mapping."""
    ctx.require(Permission.PRODUCT_IMPORT)
    checksum = validate_upload(
        filename=filename,
        content_type=content_type,
        content=content,
        allowed_types=ALLOWED_IMPORT_TYPES,
    )
    headers, rows = parse_spreadsheet(content, filename)
    if len(rows) > MAX_IMPORT_ROWS:
        raise ValidationError(
            f"This file has {len(rows)} rows; the limit is {MAX_IMPORT_ROWS}. "
            "Split it into smaller files."
        )

    storage = get_storage()
    key = storage.save(ctx.tenant_id, filename, content)
    asset = FileAsset(
        tenant_id=ctx.tenant_id,
        filename=filename,
        content_type=content_type,
        size_bytes=len(content),
        storage_backend="local",
        storage_key=key,
        checksum_sha256=checksum,
        uploaded_by_id=ctx.user_id,
    )
    db.add(asset)
    db.flush()

    mapping = suggest_mapping(headers)
    job = ImportJob(
        tenant_id=ctx.tenant_id,
        status=ImportJobStatus.UPLOADED,
        file_asset_id=asset.id,
        original_filename=filename,
        detected_headers_json=json.dumps(headers),
        column_mapping_json=json.dumps(mapping),
        branch_id=branch_id,
        total_rows=len(rows),
        created_by_id=ctx.user_id,
    )
    db.add(job)
    db.flush()
    return job, headers, mapping


def _to_decimal(value: str | None) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    text = str(value).replace(",", "").replace("ETB", "").replace("Br", "").strip()
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"'{value}' is not a number") from exc


def _row_to_fields(row: ParsedRow, mapping: dict[str, str]) -> tuple[dict, list[RowIssue]]:
    fields: dict = {}
    issues: list[RowIssue] = []
    for header, field_name in mapping.items():
        raw = row.values.get(header)
        spec = IMPORTABLE_FIELDS.get(field_name)
        if spec is None:
            continue
        if spec.get("numeric"):
            try:
                number = _to_decimal(raw)
            except ValueError as exc:
                issues.append(
                    RowIssue(row.row_number, header, "not_a_number", str(exc), raw)
                )
                continue
            if number is not None and number < 0:
                issues.append(
                    RowIssue(
                        row.row_number, header, "negative_value",
                        f"{spec['label']} cannot be negative", raw,
                    )
                )
                continue
            fields[field_name] = number
        elif raw:
            fields[field_name] = str(raw).strip()

    if not fields.get("name"):
        issues.append(
            RowIssue(row.row_number, None, "missing_name", "Product name is required", None)
        )
    return fields, issues


def validate_import(
    db: Session, ctx: AuthContext, job_id: uuid.UUID, mapping: dict[str, str] | None = None
) -> ImportPreview:
    """Validate every row and report problems before anything is written."""
    ctx.require(Permission.PRODUCT_IMPORT)
    job = get_tenant_object(db, ImportJob, job_id, ctx.tenant_id, label="Import job")
    if job.status in (ImportJobStatus.COMPLETED, ImportJobStatus.CANCELLED):
        raise ConflictError("This import has already finished")

    if mapping is not None:
        unknown = set(mapping.values()) - set(IMPORTABLE_FIELDS)
        if unknown:
            raise ValidationError(
                f"Unknown import fields: {', '.join(sorted(unknown))}",
                details={"allowed": sorted(IMPORTABLE_FIELDS)},
            )
        job.column_mapping_json = json.dumps(mapping)
    mapping = json.loads(job.column_mapping_json or "{}")
    if "name" not in mapping.values():
        raise ValidationError("Map a column to the product name before continuing")

    content = _job_content(db, ctx, job)
    _, rows = parse_spreadsheet(content, job.original_filename)

    # Clear through the relationship (delete-orphan cascade) so ``job.row_errors``
    # reflects this validation run when the commit step reads it.
    job.row_errors.clear()
    db.flush()

    existing_skus = {
        sku.lower()
        for (sku,) in db.execute(
            tenant_query(ProductVariant, ctx.tenant_id).with_only_columns(ProductVariant.sku)
        )
        if sku
    }
    existing_barcodes = {
        code.lower()
        for (code,) in db.execute(
            tenant_query(ProductVariant, ctx.tenant_id).with_only_columns(ProductVariant.barcode)
        )
        if code
    }
    existing_names = {
        name.lower()
        for (name,) in db.execute(
            tenant_query(Product, ctx.tenant_id).with_only_columns(Product.name)
        )
        if name
    }

    seen_skus: set[str] = set()
    seen_barcodes: set[str] = set()
    issues: list[RowIssue] = []
    sample: list[dict] = []
    valid = duplicates = 0

    for row in rows:
        fields, row_issues = _row_to_fields(row, mapping)
        sku = (fields.get("sku") or "").lower()
        barcode = (fields.get("barcode") or "").lower()
        name = (fields.get("name") or "").lower()

        if sku:
            if sku in existing_skus:
                row_issues.append(
                    RowIssue(row.row_number, "sku", "duplicate_sku",
                             f"SKU '{fields['sku']}' already exists in your catalogue", fields.get("sku"))
                )
            elif sku in seen_skus:
                row_issues.append(
                    RowIssue(row.row_number, "sku", "duplicate_sku_in_file",
                             f"SKU '{fields['sku']}' appears more than once in this file", fields.get("sku"))
                )
            seen_skus.add(sku)
        if barcode:
            if barcode in existing_barcodes:
                row_issues.append(
                    RowIssue(row.row_number, "barcode", "duplicate_barcode",
                             f"Barcode '{fields['barcode']}' already exists", fields.get("barcode"))
                )
            elif barcode in seen_barcodes:
                row_issues.append(
                    RowIssue(row.row_number, "barcode", "duplicate_barcode_in_file",
                             f"Barcode '{fields['barcode']}' appears more than once", fields.get("barcode"))
                )
            seen_barcodes.add(barcode)
        if name and name in existing_names and not sku and not barcode:
            duplicates += 1
            row_issues.append(
                RowIssue(row.row_number, "name", "possible_duplicate_name",
                         f"A product called '{fields['name']}' already exists", fields.get("name"))
            )

        if row_issues:
            issues.extend(row_issues)
        else:
            valid += 1
            if len(sample) < 10:
                sample.append({"row_number": row.row_number, **{k: str(v) for k, v in fields.items()}})

    for issue in issues:
        job.row_errors.append(
            ImportRowError(
                job_id=job.id,
                row_number=issue.row_number,
                column_name=issue.column,
                code=issue.code,
                message=issue.message,
                raw_value=issue.raw_value,
            )
        )

    error_row_numbers = {issue.row_number for issue in issues}
    job.total_rows = len(rows)
    job.valid_rows = valid
    job.error_rows = len(error_row_numbers)
    job.status = ImportJobStatus.VALIDATED
    db.flush()

    return ImportPreview(
        job_id=job.id,
        total_rows=len(rows),
        valid_rows=valid,
        error_rows=len(error_row_numbers),
        duplicate_rows=duplicates,
        sample=sample,
        issues=issues,
    )


def commit_import(
    db: Session,
    ctx: AuthContext,
    job_id: uuid.UUID,
    *,
    skip_rows_with_errors: bool = True,
    branch_id: uuid.UUID | None = None,
) -> ImportJob:
    """Write the valid rows.  Rows with problems are skipped and reported."""
    ctx.require(Permission.PRODUCT_IMPORT, Permission.PRODUCT_MANAGE)
    job = get_tenant_object(db, ImportJob, job_id, ctx.tenant_id, label="Import job")
    if job.status == ImportJobStatus.COMPLETED:
        raise ConflictError("This import has already been applied")
    if job.status != ImportJobStatus.VALIDATED:
        raise ConflictError("Validate the import before applying it")
    if job.error_rows and not skip_rows_with_errors:
        raise ValidationError(
            f"{job.error_rows} row(s) need attention. Fix them, or re-run allowing "
            "rows with problems to be skipped."
        )

    mapping = json.loads(job.column_mapping_json or "{}")
    content = _job_content(db, ctx, job)
    _, rows = parse_spreadsheet(content, job.original_filename)
    bad_rows = {error.row_number for error in job.row_errors}

    created = skipped = 0
    for row in rows:
        if row.row_number in bad_rows:
            skipped += 1
            continue
        fields, row_issues = _row_to_fields(row, mapping)
        if row_issues:
            skipped += 1
            continue
        create_product(
            db,
            ctx,
            ProductInput(
                name=fields["name"],
                sku=fields.get("sku"),
                description=fields.get("description"),
                brand=fields.get("brand"),
                unit_of_measure=fields.get("unit_of_measure") or "pcs",
                category_name=fields.get("category_name"),
                barcode=fields.get("barcode"),
                purchase_price=fields.get("purchase_price"),
                transport_cost=fields.get("transport_cost"),
                other_costs=fields.get("other_costs"),
                selling_price=fields.get("selling_price"),
                opening_stock=fields.get("opening_stock"),
                min_stock=fields.get("min_stock"),
                reorder_level=fields.get("reorder_level"),
                branch_id=branch_id or job.branch_id,
            ),
        )
        created += 1

    job.created_products = created
    job.skipped_rows = skipped
    job.status = ImportJobStatus.COMPLETED
    job.completed_at = utcnow()
    record_audit(
        db,
        action=AuditAction.PRODUCT_IMPORTED,
        tenant_id=ctx.tenant_id,
        actor_user_id=ctx.user_id,
        entity_type="import_job",
        entity_id=job.id,
        summary=f"Imported {created} product(s), skipped {skipped}",
        payload={"filename": job.original_filename},
    )
    db.flush()
    return job


def _job_content(db: Session, ctx: AuthContext, job: ImportJob) -> bytes:
    asset = get_tenant_object(
        db, FileAsset, job.file_asset_id, ctx.tenant_id, label="Uploaded file"
    )
    return get_storage().load(asset.storage_key)


def template_columns() -> list[dict]:
    """Columns for the downloadable import template (PRD 8.3)."""
    return [
        {"field": key, "label": spec["label"], "required": spec["required"]}
        for key, spec in IMPORTABLE_FIELDS.items()
    ]
